# -*- coding: utf-8 -*-
import io
from pptx import Presentation
from pptx.util import Emu
p = Presentation(r"H:\zpp011_v2\work\v20_8月vs9月.pptx")
s = p.slides[3]  # 第4页
texts = []
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        texts.append(sh.text_frame.text.strip().split("\n")[0])
print("slide4 texts:", texts[:8])
# 找表格
for sh in s.shapes:
    if getattr(sh, "has_table", False) and sh.has_table:
        t = sh.table
        print("table rows:", len(t.rows), "row0 h(emu):", t.rows[0].height, "row1 h(emu):", t.rows[1].height)
        print("row1 h(in):", Emu(t.rows[1].height).inches if hasattr(Emu(t.rows[1].height), "inches") else t.rows[1].height)
