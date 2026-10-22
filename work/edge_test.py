# -*- coding: utf-8 -*-
import sys, importlib.util, logging
sys.path.insert(0, r"H:\zpp011_v2\libs")
spec = importlib.util.spec_from_file_location("zpp2", r"H:\zpp011_v2\work\云南达利积压分析_v2.0_优化版.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
logging.disable(logging.WARNING)
base = r"H:\26年报表\积压表"
f1 = base + r"\材料积压--食品（10云南）202608.xlsx"
d1 = base + r"\材料积压--饮料（10云南）202608.xlsx"
f2 = base + r"\材料积压--食品（10云南）202609.xlsx"
d2 = base + r"\材料积压--饮料（10云南）202609.xlsx"
from pptx import Presentation

# case 1: 缺饮料
p1 = m.generate_ppt(f1, None, f2, None, out=r"H:\zpp011_v2\work\edge_no_bev.pptx")
print("no-bev slides:", len(Presentation(p1).slides))
# case 2: 月份反序（9月 vs 8月）→ 应自动对调
p2 = m.generate_ppt(f2, d2, f1, d1, out=r"H:\zpp011_v2\work\edge_swapped.pptx")
print("swapped slides:", len(Presentation(p2).slides))
# case 3: 全缺 → 占位
p3 = m.generate_ppt(None, None, None, None, out=r"H:\zpp011_v2\work\edge_empty.pptx")
print("empty slides:", len(Presentation(p3).slides))
print("ALL EDGE OK")
