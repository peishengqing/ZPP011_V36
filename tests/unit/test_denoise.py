# -*- coding: utf-8 -*-
"""数据层降噪（规格 §5）：汇总统计按|金额|降序整行重排 + 趋势页清杂行/空白填—。"""
import os
import sys
from openpyxl import Workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, ROOT)

from analysis.excel_builder.post_process import _resort_summary, _denoise_trend_sheet


def test_resort_summary_desc_by_abs_amount():
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总统计"
    ws.append(["序号", "车间", "总偏差金额(含税)"])
    ws.append([1, "A", -500.0])
    ws.append([2, "B", 300.0])
    ws.append([3, "C", -1200.0])
    _resort_summary(wb)
    order = [ws.cell(r, 2).value for r in (2, 3, 4)]
    assert order == ["C", "A", "B"], order
    # 数值随行搬移：C 行金额仍是 -1200（零改值）
    assert ws.cell(2, 3).value == -1200.0
    # 序号随行走（对账锚点保留原生成序）
    assert ws.cell(2, 1).value == 3


def test_resort_summary_noop_without_amount_col():
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总统计"
    ws.append(["序号", "车间"])
    ws.append([1, "A"])
    assert _resort_summary(wb) == 0  # 无金额列 → 不动


def _mk_trend():
    """按真实结构：R1 说明 / R2 表头 / R3 杂空行 / R4+ 数据。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "趋势分析（自然日分组）"
    ws.append(["早期：10/03～10/09  |  中期：10/09～10/15"])
    ws.append(["最早订单日期", "最新订单日期", "物料名称", "早期偏差率", "中期偏差率", "近期偏差率", "趋势"])
    ws.append([None] * 7)
    ws.append(["10/12", "10/17", "酱DL", None, 0.589, 1.108, "↑ 近期变差"])
    ws.append(["10/16", "10/16", "增稠剂", None, None, 1, "→"])
    return wb, ws


def test_denoise_trend_fills_blank_keeps_rows():
    """不删行（CF 区间 G3:G1086 不随行移，删行会错位）：只挪说明+填空白+重建筛选。"""
    wb, ws = _mk_trend()
    ws.merge_cells('A1:G1')   # 真实结构：R1 说明是合并单元格
    filled = _denoise_trend_sheet(wb)
    # 数据行位置不动
    assert ws.cell(4, 3).value == "酱DL"
    # 偏差率列空白填 —（趋势文本列不动）
    assert ws.cell(4, 4).value == "—"
    assert ws.cell(5, 4).value == "—" and ws.cell(5, 5).value == "—"
    assert ws.cell(4, 7).value == "↑ 近期变差"
    # 全空杂行 R3 不填（保持空）
    assert ws.cell(3, 4).value is None
    assert filled == 3
    # autofilter 覆盖数据区
    assert ws.auto_filter.ref is not None and ws.auto_filter.ref.startswith("A2")
    # R1 说明挪入表头首格批注（数据区不夹文本）
    assert ws.cell(1, 1).value in (None, "")
    assert ws.cell(2, 1).comment is not None
