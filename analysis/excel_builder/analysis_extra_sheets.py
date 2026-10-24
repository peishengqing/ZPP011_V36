# -*- coding: utf-8 -*-
"""
analysis_extra_sheets.py — 偏差分析扩展表（5 张）

从「完整偏差明细」这一张源表**一次扫描**，派生 5 张分析表：

    ① 订单级汇总          —— 一单一行：亏了多少、是不是完全没投料
    ② 车间维度            —— 工厂 × 车间：哪个车间最差
    ③ 备注归类映射        —— 原始备注写法 → 标准原因：人到底是怎么写的
    ④ 系统问题与真实超耗  —— 把「该谁负责」拆开：系统问题 / 真实超耗 / 其他
    ⑤ 物料画像            —— 一个物料一行：哪些物料是问题物料

设计约束（改本文件前先读）
------------------------------------
1. **只读一张源表**：「完整偏差明细」。5 张表共用同一次扫描的结果，**不重复扫表**
   （源表 9000+ 行 × 24 列，扫 5 遍会明显拖慢导出）。
2. **不改任何既有文件**：本模块只新增 sheet，不动 analyzer / sheet1~10 /
   post_process；表现层（数字格式 / 冻结 / 自动筛选 / 条件格式）交给 post_process
   的逐表流水线，本模块**只写数据**。
3. **数据区下方不许加说明行**：post_process 的 _last_data_row 只探前 8 列，
   说明行会被当成数据行 → 「📋 分析说明」里事实卡的「数据行数」错位 1。
   口径说明统一由 post_process 放进事实卡，不在本表里重复写。
4. **每张表各自 try**：任何一张表算挂了都不许影响其它 4 张，且必须 print 诊断
   （禁止裸 except: pass），否则导出会静默少表而没人发现。
5. **幂等**：同名 sheet 已存在先删再写，重复调用不会叠加出两张表。

最容易搞错的三处口径坑（本文件的核心防线）
------------------------------------
A. **偏差率是「比率小数」，不是百分点**
   0.143 表示 14.3%，-1 表示 -100%。写进表里的「最大偏差率绝对值」也必须是
   比率小数，post_process 才会按 0.00% 显示成 240.60%；若误当成百分点写 240.6，
   会显示成 24060.00%（放大 100 倍）。
B. **±100% 是钳位假值，不是真实超耗**
   定额=0 的行偏差率恒为 +100%（SAP 没推定额），实际=0 的行恒为 -100%。
   这两类已被「系统无定额」「完全未投料」精确覆盖；若用 >= 100 判「真实超耗」，
   会把上千行钳位行重复算成超耗，真信号被淹没。故用**严格大于 100 个百分点**
   （浮点容差 1e-6）。
C. **偏差金额是「行金额直接相减」口径**（v43.254 起）
   = 金额-实际(含税) − 金额-定额(含税)，未投料行不再被整列清零。
   本模块只**读**该列、不重算金额，避免两套算法打架。
"""
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from analysis.excel_builder.post_process import (
    _cells_of,
    _find_col,
    _header_info,
    _last_data_row,
    _num,
)
from utils.helpers import standardize_remark


# ---------------------------------------------------------------- 常量

# 唯一源表
_SRC_SHEET = '完整偏差明细'

# 5 张表的表名（**必须精确**：post_process 的事实卡按 sheet 名对账）
SHEET_ORDER = '订单级汇总'
SHEET_WORKSHOP = '车间维度'
SHEET_REMARK = '备注归类映射'
SHEET_SYSTEM = '系统问题与真实超耗'
SHEET_MATERIAL = '物料画像'

# 「实际用量为 0」的容差：与 post_process 的「完全未投料订单」表保持同一口径，
# 否则同一批订单在两张表里会给出不同的「完全未投料」结论。
_ZERO_TOL = 0.001

# 「真实超耗」的阈值：**百分点** 100（即比率小数 1.0）。
# 严格大于 + 1e-6 容差 —— ±100% 是钳位假值（定额=0 / 实际=0），必须排除。
_OVERRUN_PCT = 100.0
_OVERRUN_TOL = 1e-6

# 比率型输出的浮点精度（偏差率是比率小数，保留 6 位足够）
_RATE_ROUND = 6
# 占比型输出的精度
_RATIO_ROUND = 4


def _log(msg):
    """诊断输出：出错必须能被看见，禁止裸吞异常。"""
    print('[analysis_extra_sheets] %s' % msg)


# ---------------------------------------------------------------- 读取源表


def _id_str(v):
    """标识列统一成文本：整数去掉浮点尾巴，避免 20000387.0 这种脏值。"""
    if v is None:
        return ''
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _is_blank(v):
    """空值判定：None / 空串 / 占位符（'-' '—' 'nan' 'None'）都算空。

    与 post_process._normalize_empties 的占位符口径一致 —— 源表里把「没有值」
    写成 '-' 的行，不能被算成「有备注」。
    """
    if v is None:
        return True
    if isinstance(v, str):
        s = v.strip().lower()
        return s in ('', '-', '—', 'nan', 'none')
    return False


def _rate_ratio(v, fmt):
    """把「偏差率」单元格统一换算成**比率小数**（0.143 表示 14.3%）。

    源表有两种存储形态，都要认：
      · 文本 '19.6%'（analyzer 导出时的原始形态）→ 19.6 个百分点 → 0.196
      · 数值 + 百分比格式（post_process 规范化之后的形态）→ 本身就是比率小数 → 原值

    绝不能把它当百分点数字直接用：0.143 ≠ 14.3，差 100 倍。
    """
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        fv = float(v)
        if '%' in str(fmt or ''):
            # Excel 的百分比格式里存的就是比率小数
            return fv
        # 没有百分比格式的裸数值：按百分点处理（原始导出即如此）
        return fv / 100.0
    s = str(v).strip()
    if not s:
        return None
    if s.endswith('%'):
        s = s[:-1]
    try:
        return float(s) / 100.0
    except ValueError:
        return None


def _cell_value(cells, r, col):
    """读一个单元格的值（只取真实存在的单元格，不实体化空白格）。"""
    c = cells.get((r, col))
    return c.value if c is not None else None


def _cell_fmt(cells, r, col):
    """读一个单元格的数字格式（判断偏差率是文本百分比还是数值百分比）。"""
    c = cells.get((r, col))
    return c.number_format if c is not None else None


def _scan_source(ws):
    """一次扫描「完整偏差明细」，把明细读进内存（5 张表共用）。

    只读真实存在的单元格（_cells_of），不用 ws.cell() 去扫空白位 ——
    否则会把成千上万个空单元格实体化，保存耗时与文件体积都会上去。
    """
    hdr_row, headers = _header_info(ws)
    last_row = _last_data_row(ws, hdr_row)
    col = {
        'ord': _find_col(headers, '流程订单'),
        'date': _find_col(headers, '订单日期'),
        'fac': _find_col(headers, '工厂'),
        'ws': _find_col(headers, '车间'),
        'otype': _find_col(headers, '订单类型'),
        'pno': _find_col(headers, '产品物料号码'),
        'pdesc': _find_col(headers, '产品物料描述'),
        'mcode': _find_col(headers, '物料编码'),
        'mname': _find_col(headers, '物料名称'),
        'mtype': _find_col(headers, '物料类型'),
        'unit': _find_col(headers, '单位'),
        'quota': _find_col(headers, '定额'),
        'actual': _find_col(headers, '实际'),
        'devqty': _find_col(headers, '偏差数量'),
        'rate': _find_col(headers, '偏差率'),
        'amount': _find_col(headers, '偏差金额'),
        'note': _find_col(headers, '备注'),
        'alt': _find_col(headers, '是否替代料'),
    }
    missing = [k for k, v in col.items() if v is None]
    # 只对「所有表都要用」的两列硬失败：缺了它们，任何一张表都是无意义的垃圾结果
    # （比如缺 流程订单 → 全表挤成一个空名分组，看起来有数据其实全错）。
    # 其余列缺失只降级（例如缺 偏差率 → ④ 的「真实超耗」恒不命中，全部落「其他偏差」），
    # 降级信息必须 print 出来，不能静默出数。
    fatal = [k for k in ('ord', 'amount') if k in missing]
    if fatal:
        raise ValueError('源表缺少必要列: %s' % ', '.join(
            {'ord': '流程订单', 'amount': '偏差金额'}[k] for k in fatal))
    if missing:
        _log('源表缺少以下列，相关口径已降级: %s' % ', '.join(missing))

    cells = _cells_of(ws)
    rows = []
    for r in range(hdr_row + 1, last_row + 1):
        quota = _num(_cell_value(cells, r, col['quota']))
        actual = _num(_cell_value(cells, r, col['actual']))
        devqty = _num(_cell_value(cells, r, col['devqty']))
        amount = _num(_cell_value(cells, r, col['amount']))
        note = _cell_value(cells, r, col['note'])
        rows.append({
            'ord': _id_str(_cell_value(cells, r, col['ord'])),
            'date': _cell_value(cells, r, col['date']),
            'fac': _id_str(_cell_value(cells, r, col['fac'])),
            'ws': _id_str(_cell_value(cells, r, col['ws'])),
            'otype': _id_str(_cell_value(cells, r, col['otype'])),
            'pno': _id_str(_cell_value(cells, r, col['pno'])),
            'pdesc': _id_str(_cell_value(cells, r, col['pdesc'])),
            'mcode': _id_str(_cell_value(cells, r, col['mcode'])),
            'mname': _id_str(_cell_value(cells, r, col['mname'])),
            'mtype': _id_str(_cell_value(cells, r, col['mtype'])),
            'unit': _id_str(_cell_value(cells, r, col['unit'])),
            # 定额/实际缺失按 0 处理：与 post_process「完全未投料订单」同一口径
            'quota': 0.0 if quota is None else quota,
            'actual': 0.0 if actual is None else actual,
            'devqty': 0.0 if devqty is None else devqty,
            'amount': 0.0 if amount is None else amount,
            # 偏差率统一成比率小数（见 _rate_ratio 的说明）
            'rate': _rate_ratio(_cell_value(cells, r, col['rate']),
                                _cell_fmt(cells, r, col['rate'])),
            'note': '' if _is_blank(note) else str(note).strip(),
            'alt': _id_str(_cell_value(cells, r, col['alt'])),
        })
    return rows


# ---------------------------------------------------------------- ① 订单级汇总


def _build_order_summary(rows):
    """① 订单级汇总：一个流程订单一行。

    口径要点：
      · 订单日期/工厂/车间/订单类型/产品物料号码/产品物料描述取该订单**第一行**；
      · 投料状态按订单内**所有行**的实际用量判三态（全部为 0 / 有 0 有非 0 / 全非 0）；
      · 排序按偏差金额**绝对值**降序 —— 亏得最多的先看（亏与赚都要看绝对值，
        只按金额降序会把巨额多耗压到列表末尾）。
    """
    groups = {}
    for r in rows:
        key = r['ord']
        g = groups.get(key)
        if g is None:
            g = groups[key] = {
                'first': r,
                'rows': 0,
                'mats': set(),
                'quota': 0.0,
                'actual': 0.0,
                'devqty': 0.0,
                'amount': 0.0,
                'max_rate': 0.0,
                'no_note': 0,
                'zero_actual': 0,
                'nonzero_actual': 0,
            }
        g['rows'] += 1
        if r['mcode']:
            g['mats'].add(r['mcode'])
        g['quota'] += r['quota']
        g['actual'] += r['actual']
        g['devqty'] += r['devqty']
        g['amount'] += r['amount']
        if r['rate'] is not None:
            g['max_rate'] = max(g['max_rate'], abs(r['rate']))
        if not r['note']:
            g['no_note'] += 1
        if abs(r['actual']) <= _ZERO_TOL:
            g['zero_actual'] += 1
        else:
            g['nonzero_actual'] += 1

    out = []
    for g in groups.values():
        f = g['first']
        if g['zero_actual'] == g['rows']:
            status = '完全未投料'
        elif g['zero_actual'] > 0:
            status = '部分未投料'
        else:
            status = '正常'
        out.append([
            f['ord'],
            f['date'],
            f['fac'],
            f['ws'],
            f['otype'],
            f['pno'],
            f['pdesc'],
            g['rows'],
            len(g['mats']),
            round(g['quota'], 3),
            round(g['actual'], 3),
            round(g['devqty'], 3),
            round(g['amount'], 2),
            round(g['max_rate'], _RATE_ROUND),
            g['no_note'],
            status,
        ])
    out.sort(key=lambda x: -abs(x[12]))
    headers = ['流程订单', '订单日期', '工厂', '车间', '订单类型', '产品物料号码',
               '产品物料描述', '明细行数', '组件数(去重)', '定额数量合计',
               '实际数量合计', '偏差数量合计', '偏差金额(含税)合计',
               '最大偏差率绝对值', '无备注条数', '投料状态']
    widths = [16, 12, 18, 16, 10, 14, 30, 10, 12, 14, 14, 14, 18, 16, 12, 12]
    return headers, out, widths


# ---------------------------------------------------------------- ② 车间维度


def _workshop_type(name):
    """车间类型：含「线」→ 产线；含「车间」→ 车间；其余（综合组/配料中心…）→ 班组。"""
    if '线' in name:
        return '产线'
    if '车间' in name:
        return '车间'
    return '班组'


def _build_workshop(rows):
    """② 车间维度：按「工厂 + 车间」两个维度**一起**分组。

    口径要点：
      · 只按车间名分组会把两个厂的车间混在一起（实测「综合组」两厂都有），
        故分组键是 (工厂, 车间)；
      · 「跨厂同名」列专门标出这个坑：同一车间名出现在多于一个工厂时标「是」；
      · 备注覆盖率 = 有备注条数 ÷ 明细行数（比率小数，明细行数为 0 时给 0）。
    """
    factories = {}          # 车间名 -> 该车间名出现过的工厂集合
    groups = {}
    for r in rows:
        factories.setdefault(r['ws'], set()).add(r['fac'])
        key = (r['fac'], r['ws'])
        g = groups.get(key)
        if g is None:
            g = groups[key] = {
                'rows': 0, 'dev_rows': 0, 'devqty': 0.0, 'amount': 0.0,
                'has_note': 0, 'no_note': 0,
            }
        g['rows'] += 1
        if abs(r['devqty']) > 1e-9:
            g['dev_rows'] += 1
        g['devqty'] += r['devqty']
        g['amount'] += r['amount']
        if r['note']:
            g['has_note'] += 1
        else:
            g['no_note'] += 1

    out = []
    for (fac, ws), g in groups.items():
        coverage = (g['has_note'] / g['rows']) if g['rows'] else 0.0
        out.append([
            fac,
            ws,
            _workshop_type(ws),
            g['rows'],
            g['dev_rows'],
            round(g['devqty'], 3),
            round(g['amount'], 2),
            g['has_note'],
            g['no_note'],
            round(coverage, _RATIO_ROUND),
            '是' if len(factories.get(ws, ())) > 1 else '否',
        ])
    out.sort(key=lambda x: -abs(x[6]))
    headers = ['工厂', '车间', '车间类型', '明细行数', '偏差条数', '偏差数量合计',
               '偏差金额(含税)合计', '有备注条数', '无备注条数', '备注覆盖率', '跨厂同名']
    widths = [18, 20, 10, 10, 10, 14, 18, 12, 12, 12, 10]
    return headers, out, widths


# ---------------------------------------------------------------- ③ 备注归类映射


def _build_remark_map(rows):
    """③ 备注归类映射：按**原始写法**分组，再映射到标准原因。

    口径要点：
      · 分组键是原始写法（去首尾空白后），不是标准原因 —— 目的正是暴露
        「同一个意思写了七八种」（系统无定额 / 系统无额定 / 系统未维护…）；
      · 空备注不进本表（分母「全部非空备注条数」也因此只算非空的）；
      · 标准原因统一由 utils.helpers.standardize_remark 产出，不在本模块里
        另写一套关键词表（两处词表打架是这类报表最常见的口径事故）。
    """
    groups = {}
    total = 0
    for r in rows:
        raw = r['note']
        if not raw:
            continue
        total += 1
        g = groups.get(raw)
        if g is None:
            g = groups[raw] = {'n': 0, 'orders': set(), 'amount': 0.0}
        g['n'] += 1
        g['orders'].add(r['ord'])
        g['amount'] += r['amount']

    out = []
    for raw, g in groups.items():
        ratio = (g['n'] / total) if total else 0.0
        out.append([
            raw,
            standardize_remark(raw),
            g['n'],
            len(g['orders']),
            round(g['amount'], 2),
            round(ratio, _RATIO_ROUND),
        ])
    out.sort(key=lambda x: -x[2])
    headers = ['原始备注写法', '标准原因', '明细条数', '涉及订单数',
               '偏差金额(含税)合计', '占非空备注比']
    widths = [46, 16, 10, 12, 18, 14]
    return headers, out, widths


# ---------------------------------------------------------------- ④ 系统问题与真实超耗


def _classify(row):
    """行级分类（优先级从上到下，命中即止）。

      1. 系统无定额：定额=0 且 实际>0 —— SAP 没推定额，是**系统问题**，不追人；
      2. 完全未投料：实际=0 且 定额>0 —— 产量报了但一克没投，是**最严重**的问题；
      3. 真实超耗：|偏差率| 严格超过 100 个百分点，且不属于上面两类
         （±100% 是钳位假值，不是真超耗）；
      4. 其他偏差：其余全部。
    """
    if row['quota'] == 0 and row['actual'] > 0:
        return '系统无定额'
    if abs(row['actual']) <= _ZERO_TOL and row['quota'] > 0:
        return '完全未投料'
    rate = row['rate']
    if rate is not None and abs(rate) * 100.0 > _OVERRUN_PCT + _OVERRUN_TOL:
        return '真实超耗'
    return '其他偏差'


def _build_system_issue(rows):
    """④ 系统问题与真实超耗：行级分类后按类汇总。

    口径要点：
      · 「占全期偏差金额绝对值比」= 该类**行级金额绝对值之和** ÷ 全部行金额绝对值之和
        （分母为 0 时给 0），四类相加 = 100%（允许四舍五入误差）；
      · 占比按**行级绝对值**算（不是 |类合计|），因为同一类内可能有正负相抵
        （如「其他偏差」里 +100 万与 -100 万互抵），用类合计会漏掉这部分、
        导致四类加和 ≠ 100%；
      · 四类**全部输出**（哪怕某类 0 行），表格结构稳定、便于逐期对比。

    ⚠️ **分母不能用「净额」**：净额是多耗与少耗互相抵消后的残差，本期多耗与少耗
      几乎抵消（净额仅几百万），单类占比会算出 392% / -394% 这种没法读的数。
      改用行级绝对值之和做分母后，占比恒在 0~100% 且四类加和为 1，
      读作「该类在全部偏差金额里占多大分量」。**判断正负方向请看金额列的符号。**
    """
    cats = ['系统无定额', '完全未投料', '真实超耗', '其他偏差']
    agg = {c: {'rows': 0, 'orders': set(), 'devqty': 0.0, 'amount': 0.0,
               'abs_amount': 0.0} for c in cats}
    abs_total = 0.0
    for r in rows:
        cat = _classify(r)
        g = agg[cat]
        g['rows'] += 1
        g['orders'].add(r['ord'])
        g['devqty'] += r['devqty']
        g['amount'] += r['amount']
        g['abs_amount'] += abs(r['amount'])
        abs_total += abs(r['amount'])

    out = []
    for c in cats:
        g = agg[c]
        share = (g['abs_amount'] / abs_total) if abs_total else 0.0
        out.append([
            c,
            g['rows'],
            len(g['orders']),
            round(g['devqty'], 3),
            round(g['amount'], 2),
            round(share, _RATIO_ROUND),
        ])
    out.sort(key=lambda x: -x[1])
    headers = ['分类', '明细行数', '涉及订单数', '偏差数量合计',
               '偏差金额(含税)合计', '占全期偏差金额绝对值比']
    widths = [16, 10, 12, 14, 18, 16]
    return headers, out, widths


# ---------------------------------------------------------------- ⑤ 物料画像


def _build_material(rows):
    """⑤ 物料画像：一个物料编码一行。

    口径要点：
      · 物料名称/物料类型/单位取该物料**第一行**的值（同料不同行的描述以首次出现为准）；
      · 含替代料 = 该物料**任一行**是否替代料 == '是'（物料级标记，不是行级）；
      · 备注覆盖率 = 该物料有备注行数 ÷ 该物料总行数。
    """
    groups = {}
    for r in rows:
        key = r['mcode']
        g = groups.get(key)
        if g is None:
            g = groups[key] = {
                'first': r, 'rows': 0, 'orders': set(), 'dev_rows': 0,
                'devqty': 0.0, 'amount': 0.0, 'max_rate': 0.0,
                'has_note': 0, 'alt': False,
            }
        g['rows'] += 1
        g['orders'].add(r['ord'])
        if abs(r['devqty']) > 1e-9:
            g['dev_rows'] += 1
        g['devqty'] += r['devqty']
        g['amount'] += r['amount']
        if r['rate'] is not None:
            g['max_rate'] = max(g['max_rate'], abs(r['rate']))
        if r['note']:
            g['has_note'] += 1
        if r['alt'] == '是':
            g['alt'] = True

    out = []
    for g in groups.values():
        f = g['first']
        coverage = (g['has_note'] / g['rows']) if g['rows'] else 0.0
        out.append([
            f['mcode'],
            f['mname'],
            f['mtype'],
            f['unit'],
            len(g['orders']),
            g['dev_rows'],
            round(g['devqty'], 3),
            round(g['amount'], 2),
            round(g['max_rate'], _RATE_ROUND),
            round(coverage, _RATIO_ROUND),
            '是' if g['alt'] else '否',
        ])
    out.sort(key=lambda x: -abs(x[7]))
    headers = ['物料编码', '物料名称', '物料类型', '单位', '涉及订单数', '偏差条数',
               '偏差数量合计', '偏差金额(含税)合计', '最大偏差率绝对值',
               '备注覆盖率', '含替代料']
    widths = [14, 30, 10, 8, 12, 10, 14, 18, 16, 12, 10]
    return headers, out, widths


# ---------------------------------------------------------------- 写表


def _write_sheet(wb, name, headers, data, widths):
    """写一张新表（幂等：同名先删）。返回写入的数据行数。"""
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for j in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=j)
        c.font = Font(bold=True, size=11)
        c.alignment = Alignment(horizontal='center', vertical='center')
    for row in data:
        ws.append(row)
    for j, w in enumerate(widths, 1):
        if w:
            ws.column_dimensions[get_column_letter(j)].width = w
    return len(data)


# ---------------------------------------------------------------- 入口


def build_extra_sheets(wb):
    """生成 5 张偏差分析扩展表。返回 {表名: 写入行数}；源表缺失时返回 {}。

    任何一张表失败都不影响其它表（逐表 try/except + print 诊断），
    整体绝不抛异常 —— 附加分析表出问题不该让整个导出失败。
    """
    result = {}
    if _SRC_SHEET not in wb.sheetnames:
        _log('源表「%s」不存在，跳过 5 张扩展表' % _SRC_SHEET)
        return result

    try:
        rows = _scan_source(wb[_SRC_SHEET])
    except Exception as e:
        _log('扫描源表失败，5 张扩展表全部跳过: %s' % e)
        return result

    builders = (
        (SHEET_ORDER, _build_order_summary),
        (SHEET_WORKSHOP, _build_workshop),
        (SHEET_REMARK, _build_remark_map),
        (SHEET_SYSTEM, _build_system_issue),
        (SHEET_MATERIAL, _build_material),
    )
    for name, builder in builders:
        try:
            headers, data, widths = builder(rows)
            result[name] = _write_sheet(wb, name, headers, data, widths)
        except Exception as e:
            _log('生成「%s」失败，已跳过该表: %s' % (name, e))
    return result
